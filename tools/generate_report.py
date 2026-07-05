"""基于教师 Word 模板生成课程设计报告（保留封面、目录、页眉页脚和样式）。"""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "tmp" / "report_template" / "template.docx"
OUTPUT = ROOT / "docs" / "《数字图像处理与机器视觉课程设计报告_人脸图像预处理及人脸检测系统设计》.docx"
ASSETS = ROOT / "results" / "report_assets"
METRICS = ROOT / "results" / "metrics" / "batch_metrics.json"


def set_run_font(run, chinese="宋体", western="Times New Roman", size=12, bold=False):
    run.font.name = western
    run.font.size = Pt(size)
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), chinese)
    rfonts.set(qn("w:ascii"), western)
    rfonts.set(qn("w:hAnsi"), western)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level):
    style_name = f"Heading {level}"
    if style_name not in [style.name for style in doc.styles]:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 2" if level > 2 else "Normal"]
    p = doc.add_paragraph(text, style=style_name)
    for run in p.runs:
        set_run_font(run, chinese="黑体", western="Times New Roman", size=16 if level == 1 else 14, bold=True)
    p.paragraph_format.keep_with_next = True
    return p


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, chinese="宋体", western="Cambria Math", size=11)
    r.font.italic = True


def add_code(doc, code, caption):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.right_indent = Pt(18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shading)
    r = p.add_run(code)
    set_run_font(r, chinese="等线", western="Consolas", size=8.5)
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    for run in cap.runs:
        set_run_font(run, size=10)


def configure_table(table, widths):
    table.autofit = False
    table.alignment = 1
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[i]))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15


def add_table(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    for run in cap.runs:
        set_run_font(run, size=10.5, bold=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    configure_table(table, widths)
    for ri, row in enumerate(table.rows):
        if ri == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cell in row.cells:
            if ri == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "D9EAF7")
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=9, bold=(ri == 0))
    return table


def add_figure(doc, path, caption, width=5.8):
    path = Path(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    for run in cap.runs:
        set_run_font(run, size=10.5)


def make_flowchart(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1800, 420), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    font = ImageFont.truetype(str(font_path), 36) if font_path.exists() else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 28) if font_path.exists() else ImageFont.load_default()
    labels = ["图像输入", "尺寸归一化", "灰度化", "增强/滤波", "Haar人脸检测", "ROI人眼检测", "结果与指标"]
    colors = ["#E8F2FF", "#EAF7EF", "#FFF4D6", "#FDECEB", "#EFE9FF", "#E7F6F8", "#F3F3F3"]
    x_positions = [30, 275, 520, 765, 1010, 1255, 1500]
    for i, (x, label) in enumerate(zip(x_positions, labels)):
        box = (x, 130, x + 205, 280)
        draw.rounded_rectangle(box, radius=22, fill=colors[i], outline="#3E5C76", width=4)
        bbox = draw.textbbox((0, 0), label, font=small)
        tx = x + (205 - (bbox[2] - bbox[0])) / 2
        ty = 205 - (bbox[3] - bbox[1]) / 2
        draw.text((tx, ty), label, fill="#1F2933", font=small)
        if i < len(labels) - 1:
            draw.line((x + 205, 205, x + 240, 205), fill="#246BFD", width=6)
            draw.polygon([(x + 240, 205), (x + 222, 194), (x + 222, 216)], fill="#246BFD")
    title = "人脸图像预处理及检测系统总体流程"
    bbox = draw.textbbox((0, 0), title, font=font)
    draw.text(((1800 - (bbox[2] - bbox[0])) / 2, 35), title, fill="#16324F", font=font)
    image.save(path)


def clear_template_body(doc):
    # 保留封面、目录、分节符和目录域，仅删除第一个正文一级标题之后的模板占位内容。
    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "1 课程设计概述")
    for p in list(doc.paragraphs[start:]):
        p._element.getparent().remove(p._element)


def update_cover(doc):
    doc.paragraphs[6].text = "题  目：基于 Python + OpenCV 的人脸图像"
    doc.paragraphs[7].text = "        预处理及人脸检测系统设计"
    doc.paragraphs[10].text = "日期：   2026  年   7   月   5  日"
    for index in (6, 7, 10):
        for run in doc.paragraphs[index].runs:
            set_run_font(run, chinese="宋体", size=15 if index != 10 else 12)


def build_report():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"模板 DOCX 不存在：{TEMPLATE}")
    payload = json.loads(METRICS.read_text(encoding="utf-8"))
    records = payload["records"]
    search = payload.get("parameter_search", {})
    doc = Document(TEMPLATE)
    update_cover(doc)
    clear_template_body(doc)

    flowchart = ASSETS / "system_flowchart.png"
    make_flowchart(flowchart)

    add_heading(doc, "1 课程设计概述", 1)
    add_heading(doc, "1.1 课程设计背景", 2)
    add_body(doc, "人脸检测是计算机视觉系统中的基础环节，广泛应用于人机交互、门禁安防、身份认证、智能监控和图像检索等场景。检测模块需要先从整幅图像中确定人脸位置，后续才能进行关键点定位、特征提取或身份识别。随着图像采集设备的普及，输入图像不再都是光照均匀、背景简单的标准证件照，弱光、噪声、拍摄距离变化、复杂纹理以及轻度侧脸都会影响传统检测算法的稳定性。")
    add_body(doc, "Haar 级联分类器具有结构清晰、CPU 运行速度快、工程部署简单等特点，适合作为课程设计中理解滑动窗口、特征分类和参数权衡的实验对象。但 Haar 正面模型依赖局部明暗结构，对曝光不足、噪声和大角度姿态比较敏感。因此，本设计在检测前增加灰度化、光照均衡、Gamma 校正、滤波去噪和尺寸归一化，并采用多预处理分支对比与参数搜索，观察图像质量变化对检测结果的实际影响。")

    add_heading(doc, "1.2 设计目的", 2)
    purposes = [
        "巩固图像灰度化、增强、滤波和尺寸归一化等数字图像处理基础知识。",
        "掌握 Python、NumPy 与 OpenCV 的图像处理流程，理解数组运算与图像数据之间的关系。",
        "理解 Haar 特征级联分类器完成人脸和人眼检测的基本流程及适用边界。",
        "完成包含 PyQt6 GUI、命令行、结果保存、批量处理和参数调节的完整系统。",
        "通过控制变量实验培养结果对比、参数调试和工程问题分析能力。",
    ]
    for i, text in enumerate(purposes, 1):
        add_body(doc, f"{i}、{text}")

    add_heading(doc, "1.3 主要设计内容与技术指标", 2)
    rows = [
        ["图像输入", "JPG、JPEG、PNG、BMP、TIF、TIFF", "支持单图与目录批处理"],
        ["预处理", "灰度、均衡化、CLAHE、Gamma、高斯/中值滤波", "核心预处理由 NumPy 自主实现"],
        ["尺寸归一化", "640×480", "等比例缩放并填充，避免几何拉伸"],
        ["目标检测", "Haar 正面人脸与人眼", "人眼限定在人脸上部 ROI"],
        ["参数控制", "scaleFactor、minNeighbors、minSize、Gamma", "GUI 调节并支持命令行传入"],
        ["结果输出", "JPG/PNG、CSV、JSON", "包含画框图、对比图、数量与耗时"],
        ["交互方式", "PyQt6 GUI + 命令行", "支持保存、批处理和参数搜索"],
    ]
    add_table(doc, "表 1-1 系统主要设计内容与技术指标", ["模块", "技术指标", "实现说明"], rows, [1800, 3300, 3300])

    add_heading(doc, "1.4 实验环境", 2)
    env_rows = [
        ["操作系统", "Windows 10/11（本次运行环境为 Windows）"],
        ["开发语言", "Python 3.11.15"],
        ["主要依赖", "NumPy 2.2.6、opencv-python 4.11.0.86、PyQt6 6.11.0、pytest 9.1.1"],
        ["开发与运行工具", "VS Code / PowerShell / Conda"],
        ["检测模型", "haarcascade_frontalface_default.xml、haarcascade_eye_tree_eyeglasses.xml"],
        ["计算方式", "CPU，未使用 GPU"],
    ]
    add_table(doc, "表 1-2 实验软硬件环境", ["项目", "配置"], env_rows, [2200, 6200])

    add_heading(doc, "2 相关技术与理论基础", 1)
    add_heading(doc, "2.1 彩色图像、灰度图像与灰度化", 2)
    add_body(doc, "彩色图像通常由多个颜色通道组成。OpenCV 读取彩色图像时使用 BGR 通道顺序，而显示与常见公式通常采用 RGB 顺序。灰度图像只保留亮度信息，每个像素由 0～255 的单一数值表示。Haar 分类器主要依赖局部明暗差异，因此转换为灰度图可以降低计算量。项目没有直接调用颜色转换接口，而是按亮度权重自主计算：")
    add_formula(doc, "Gray = 0.299R + 0.587G + 0.114B")
    add_body(doc, "绿色通道权重最大，是因为人眼对绿色亮度变化更敏感。计算时先把 B、G、R 通道转换为浮点数，完成加权后四舍五入并限制到 0～255。")

    add_heading(doc, "2.2 直方图均衡化与 CLAHE", 2)
    add_body(doc, "灰度直方图统计每个灰度级在图像中的出现次数。全局直方图均衡化利用累计分布函数重新映射灰度，使有效动态范围得到扩展。设灰度级总数为 L，灰度 r_k 的累计概率为 CDF(r_k)，映射关系为：")
    add_formula(doc, "s_k = round[(L - 1) × CDF(r_k)]")
    add_body(doc, "全局均衡化适合整体偏暗或对比度不足的图像，但可能把噪声和背景纹理一起放大。CLAHE 将图像划分为若干小区域，对各区域进行受限对比度均衡，并在区域之间插值。项目使用 OpenCV 的 CLAHE 基础接口，默认 clipLimit 为 2.0、网格为 8×8。")

    add_heading(doc, "2.3 Gamma 校正", 2)
    add_body(doc, "Gamma 校正是一种幂律灰度变换。项目约定 Gamma 大于 1 时抬升暗部，公式为：")
    add_formula(doc, "s = 255 × (r / 255)^(1 / gamma)")
    add_body(doc, "为了提高速度，程序预先生成 256 项查找表，再按原像素值索引映射。Gamma 可以改善弱光图像的可见性，但不能恢复传感器已经丢失的细节，取值过大还会造成亮部过曝。")

    add_heading(doc, "2.4 高斯滤波与中值滤波", 2)
    add_body(doc, "高斯滤波根据像素与中心点的距离分配权重，二维高斯函数为：")
    add_formula(doc, "G(x,y) = 1/(2πσ²) × exp[-(x²+y²)/(2σ²)]")
    add_body(doc, "项目自主生成高斯核并归一化，使核元素之和为 1，再通过自定义滑动窗口卷积完成滤波。高斯滤波适合随机噪声，但也会平滑边缘。中值滤波对邻域像素排序并取中位数，对椒盐噪声抑制效果较好，同时比线性平均更能保留边缘。")

    add_heading(doc, "2.5 尺寸归一化", 2)
    add_body(doc, "不同输入图像的分辨率和宽高比不一致，直接拉伸会改变人脸几何比例。系统计算目标宽高与原图宽高的缩放比例最小值，使用 cv2.resize 等比例缩放，再在周围填充黑边，最终统一到 640×480。缩小时使用区域插值，放大时使用三次插值。")

    add_heading(doc, "2.6 Haar 特征分类器与检测参数", 2)
    add_body(doc, "Haar 特征用相邻矩形区域的像素和之差描述眼睛、鼻梁和面颊等结构的明暗关系。积分图可以用四个点快速计算任意矩形像素和；AdaBoost 从大量候选特征中选择区分能力较强的弱分类器并组合；级联结构先用简单层快速淘汰背景窗口，再逐级验证困难候选。项目加载 OpenCV 已训练的 XML 模型，不从零训练分类器，课程设计重点放在预处理、检测流程封装、参数优化与系统集成。")
    parameter_rows = [
        ["scaleFactor", "相邻尺度的缩放比例", "越接近 1 搜索越细、耗时越高；过大可能漏过合适尺度"],
        ["minNeighbors", "候选框所需邻近矩形数量", "较小提高召回但可能误检，较大更严格但可能漏检"],
        ["minSize", "最小检测窗口", "减小有利于小脸，同时增加搜索量和误检风险"],
    ]
    add_table(doc, "表 2-1 detectMultiScale 主要参数含义", ["参数", "含义", "影响"], parameter_rows, [1800, 2600, 4000])

    add_heading(doc, "3 核心算法原理与实现", 1)
    add_heading(doc, "3.1 模块详细设计", 2)
    module_rows = [
        ["配置模块", "src/config.py", "管理预处理、Haar 参数、分类器路径和支持格式"],
        ["图像读写", "src/image_io.py", "中文路径安全读写、保存和目录扫描"],
        ["图像预处理", "src/preprocess.py", "自主灰度、Gamma、均衡、卷积、滤波及组合分支"],
        ["检测模块", "src/detector.py", "加载 Haar 模型、检测人脸/人眼并搜索参数"],
        ["评价模块", "src/evaluator.py", "统计数量、耗时、稳定性并输出 CSV/JSON"],
        ["可视化", "src/visualizer.py", "绘制检测框、生成多宫格对比图"],
        ["图形界面", "gui_app.py", "提供图像显示、按钮操作、参数控件和批处理"],
        ["命令行入口", "main.py", "解析参数并组织单图/批量处理流程"],
    ]
    add_table(doc, "表 3-1 系统模块划分", ["模块名称", "文件", "功能说明"], module_rows, [1600, 2500, 4300])
    flow_intro = add_body(doc, "系统采用分层设计。GUI 和命令行仅负责接收用户输入并组织流程，算法实现集中在 src 目录，避免界面文件堆积图像处理逻辑。总体流程如图 3-1 所示。")
    flow_intro.paragraph_format.keep_together = True
    flow_intro.paragraph_format.keep_with_next = True
    add_figure(doc, flowchart, "图 3-1 人脸图像预处理及检测系统总体流程", width=6.3)

    add_heading(doc, "3.2 核心算法详细实现", 2)
    add_heading(doc, "3.2.1 灰度化与 Gamma 校正", 3)
    add_body(doc, "灰度化显式拆分 BGR 通道并按亮度权重计算；Gamma 校正构造查找表后直接用像素值索引。关键代码如下：")
    add_code(doc, """b, g, r = image[..., 0], image[..., 1], image[..., 2]
gray = (0.114 * b.astype(np.float32)
        + 0.587 * g.astype(np.float32)
        + 0.299 * r.astype(np.float32))
gray = np.clip(np.rint(gray), 0, 255).astype(np.uint8)

table = np.array([
    ((i / 255.0) ** (1.0 / gamma)) * 255
    for i in range(256)
], dtype=np.uint8)
result = table[image]""", "代码 3-1 灰度化与 Gamma 查找表核心实现")

    add_heading(doc, "3.2.2 直方图均衡化", 3)
    add_body(doc, "程序使用 bincount 自主统计 256 级直方图，计算累计分布后排除首个非零累计值，避免纯色背景影响映射范围。")
    add_code(doc, """histogram = np.bincount(gray.ravel(), minlength=256)
cdf = histogram.cumsum()
nonzero = np.flatnonzero(cdf)
cdf_min = cdf[nonzero[0]]
denominator = gray.size - cdf_min
if denominator <= 0:
    return gray.copy()
mapping = np.clip(
    np.rint((cdf - cdf_min) * 255.0 / denominator),
    0, 255
).astype(np.uint8)
result = mapping[gray]""", "代码 3-2 手工直方图均衡化核心实现")

    add_heading(doc, "3.2.3 高斯核、自定义卷积与中值滤波", 3)
    add_body(doc, "高斯核根据坐标网格和标准差计算并归一化；卷积使用边缘复制填充和滑动窗口，既支持灰度图也支持彩色图。")
    add_code(doc, """coordinates = np.arange(-radius, radius + 1)
xx, yy = np.meshgrid(coordinates, coordinates)
kernel = np.exp(-(xx * xx + yy * yy)
                / (2.0 * sigma * sigma))
kernel = kernel / kernel.sum()

padded = np.pad(image.astype(np.float64),
                pad_width, mode=\"edge\")
windows = np.lib.stride_tricks.sliding_window_view(
    padded, kernel.shape, axis=(0, 1))
result = np.sum(windows * kernel, axis=(-2, -1))""", "代码 3-3 高斯核与自定义二维卷积")
    add_body(doc, "中值滤波同样构造滑动窗口，但不进行加权求和，而是在窗口维度取中位数，从而抑制孤立的极亮或极暗脉冲点。")
    add_code(doc, """pad = kernel_size // 2
padded = np.pad(image, pad_width, mode=\"edge\")
windows = np.lib.stride_tricks.sliding_window_view(
    padded, (kernel_size, kernel_size), axis=(0, 1))
result = np.median(windows, axis=(-2, -1))
result = np.clip(result, 0, 255).astype(image.dtype)""", "代码 3-4 中值滤波核心实现")

    add_heading(doc, "3.2.4 Haar 人脸与人眼检测", 3)
    add_body(doc, "检测器先在整幅灰度图上寻找正面人脸，再截取每张人脸上部 65% 作为眼睛 ROI。眼睛候选按面积降序排列，每张脸最多保留两个，从而减少鼻孔、嘴部和背景纹理误检。")
    add_code(doc, """faces = face_cascade.detectMultiScale(
    gray, scaleFactor=cfg.scale_factor,
    minNeighbors=cfg.min_neighbors,
    minSize=cfg.min_size)
for x, y, w, h in faces:
    roi = gray[y:y + int(h * 0.65), x:x + w]
    eyes = eye_cascade.detectMultiScale(
        roi, scaleFactor=cfg.eye_scale_factor,
        minNeighbors=cfg.eye_min_neighbors,
        minSize=cfg.eye_min_size)
    candidates = sorted(eyes, key=lambda b: b[2]*b[3], reverse=True)
    eye_boxes.extend(candidates[:2])""", "代码 3-5 人脸 ROI 内的人眼检测流程")

    add_heading(doc, "3.2.5 GUI 调用流程", 3)
    add_body(doc, "PyQt6 界面负责打开文件、显示图像和读取参数控件。一键完整处理依次调用 normalize_size、to_gray、gamma_correct、gaussian_filter、clahe_equalize 和 HaarFaceEyeDetector.detect，再将结果交给 visualizer 绘框。界面不重复实现算法，降低了 GUI 与算法模块耦合。")

    add_heading(doc, "3.3 算法优化方案", 2)
    optimizations = [
        ("弱光图像", "采用 Gamma 校正抬升暗部，再使用 CLAHE 改善局部对比度，避免只做全局拉伸。"),
        ("噪声图像", "对高斯随机噪声优先使用高斯滤波，对脉冲噪声使用中值滤波，然后再做局部增强。"),
        ("尺寸不一致", "等比例缩放并填充到 640×480，保持人脸宽高比例并统一检测尺度。"),
        ("检测参数", "遍历五组 scaleFactor、minNeighbors 和 minSize，综合非零检测比例、数量稳定性、眼睛证据和耗时选择候选。"),
        ("人眼误检", "仅在人脸上部 65% ROI 内检测，并限制每张脸最多两个大候选。"),
        ("分支对比", "并行比较原图、灰度、均衡化、CLAHE、滤波+CLAHE 和 Gamma+CLAHE，避免预处理越多越好的错误假设。"),
    ]
    for name, desc in optimizations:
        add_body(doc, f"{name}：{desc}", bold_prefix=f"{name}：")
    add_body(doc, "参数搜索使用无标注代理分数，只能帮助缩小候选范围。如果所有分支都稳定误检，代理分数仍可能偏高，因此必须结合检测框人工核验。")

    add_heading(doc, "4 实验结果分析", 1)
    add_heading(doc, "4.1 实验方案", 2)
    add_body(doc, "实验在同一台计算机上运行 python -m pytest、python main.py --help 和带 --optimize 的批处理命令。测试集包含三张 512×512 单人图像：sample_lena.jpg 为 OpenCV 官方正常光照示例；sample_lowlight.jpg 由基准图像降低亮度生成；sample_noise.jpg 由基准图像加入固定随机种子的高斯噪声生成。三张图使用同一主体，便于控制人物和构图变量。")
    add_body(doc, "对每张图分别比较 original、gray、hist_eq、clahe、gaussian_clahe、median_clahe 和 gamma_clahe 七个分支，并记录 face_count、eye_count 与 elapsed_ms。另对五组 Haar 参数进行搜索，查看检测数量、跨分支一致性、主观视觉效果和耗时。由于没有人工标注框，实验不报告虚构准确率；agreement_ratio 只表示各分支检测数量的一致程度。")

    add_heading(doc, "4.2 实验结果分析", 2)
    summary_rows = []
    for image_name in ("sample_lena.jpg", "sample_lowlight.jpg", "sample_noise.jpg"):
        image_records = [r for r in records if r["image"] == image_name]
        for r in image_records:
            method_cn = {
                "original": "归一化原图", "gray": "灰度化", "hist_eq": "直方图均衡",
                "clahe": "CLAHE", "gaussian_clahe": "高斯+CLAHE",
                "median_clahe": "中值+CLAHE", "gamma_clahe": "Gamma+CLAHE",
            }[r["method"]]
            analysis = "人脸稳定检出" if r["face_count"] == 1 else ("出现额外候选，需人工核验" if r["face_count"] > 1 else "未检出人脸")
            summary_rows.append([image_name.replace("sample_", ""), method_cn, r["face_count"], r["eye_count"], f"{r['elapsed_ms']:.3f}", analysis])
    add_table(doc, "表 4-1 不同图像与预处理分支的真实检测结果", ["图像", "预处理", "人脸", "人眼", "耗时/ms", "结果分析"], summary_rows, [1250, 1600, 800, 800, 1200, 2750])

    add_body(doc, "从表 4-1 可见，本次参数搜索完成后，三张样例的七个分支均检测到 1 张人脸，跨分支 agreement_ratio 为 1.0、人脸数量标准差为 0。该结果说明当前单人演示集上的数量稳定，但不等价于 100% 准确率。正常光照图和弱光图的大多数分支均检出 2 只眼睛；加噪图在未滤波分支可能只保留 1 只眼睛，而高斯+CLAHE 分支检出 2 只眼睛，说明先抑制噪声有助于恢复局部眼部结构。")

    selected_rows = []
    for image_name, info in search.items():
        cfg = info["selected"]
        best_trial = max(info["trials"], key=lambda t: t["score"])
        selected_rows.append([image_name, cfg["scale_factor"], cfg["min_neighbors"], f"{cfg['min_size'][0]}×{cfg['min_size'][1]}", f"{best_trial['mean_ms']:.3f}", f"{best_trial['score']:.4f}"])
    add_table(doc, "表 4-2 各样例参数搜索的选中组合", ["图像", "scaleFactor", "minNeighbors", "minSize", "平均耗时/ms", "代理分数"], selected_rows, [1900, 1200, 1200, 1200, 1500, 1400])
    add_body(doc, "三张单人大脸样例均倾向选择较大的 scaleFactor 和 minSize，因为人脸尺寸充足，稀疏尺度搜索仍能稳定检出并节省时间。该结论不能直接用于多人小脸场景；若新增远距离人物，应减小 minSize 并重新搜索。")

    figures = [
        (ROOT / "results/comparison/sample_lena_comparison.jpg", "图 4-1 正常光照图像多预处理分支的人脸与人眼检测结果"),
        (ROOT / "results/comparison/sample_lowlight_comparison.jpg", "图 4-2 弱光图像在均衡、滤波和 Gamma 处理下的检测对比"),
        (ROOT / "results/comparison/sample_noise_comparison.jpg", "图 4-3 加噪图像滤波前后的检测结果对比"),
        (ROOT / "results/gui_preview.png", "图 4-4 PyQt6 人脸图像预处理及检测系统运行界面"),
    ]
    for path, caption in figures:
        if path.exists():
            add_figure(doc, path, caption, width=6.0 if "gui" not in path.name else 6.3)
    add_body(doc, "图 4-1 至图 4-3 的绿色矩形为人脸框，红色矩形为人眼框。正常光照图各分支结果接近；弱光图经均衡后面部层次更明显，但全局均衡也增强了背景；噪声图在高斯或中值滤波后颗粒被抑制，人眼候选更完整。图 4-4 展示了 GUI 的原图/结果双栏、12 个功能按钮、四项参数控件以及检测信息区。")
    add_body(doc, "本次 pytest 共收集 7 个测试并全部通过，覆盖自主预处理算法、中文路径图像读写、Haar 模型加载、可视化、指标导出、GUI 模块导入和命令行帮助启动。命令行批处理成功处理 3 张图片，并生成 detection、comparison 和 metrics 三类结果。")

    add_heading(doc, "5 遇到的问题及解决方案", 1)
    problems = [
        ("弱光图像检测不稳定", "输入亮度低时眼眶、鼻梁和面颊之间的明暗差异减弱。", "先使用 Gamma 抬升暗部，再用 CLAHE 增强局部对比度，并避免 Gamma 过大造成亮部饱和。", "弱光样例各分支均检出 1 张人脸，增强后的面部细节更便于观察。"),
        ("噪声影响边缘和纹理特征", "随机噪声改变局部矩形区域的灰度和，干扰 Haar 特征及眼睛候选。", "按噪声类型使用高斯或中值滤波，再进行 CLAHE；同时保留未滤波分支用于对照。", "加噪样例在高斯+CLAHE 后检出 2 只眼睛，优于部分未滤波分支。"),
        ("Haar 对侧脸适应性有限", "使用的 XML 主要以正面人脸训练，大角度偏转改变了特征排列。", "对轻度偏转通过减小 minNeighbors 或 minSize 提高召回；对大角度侧脸明确作为模型边界，后续可增加 profileface 或 DNN。", "系统能处理正面和轻度偏转样例，但不宣称解决所有侧脸。"),
        ("参数过松误检、过严漏检", "较低 minNeighbors 和较小 minSize 会保留更多背景候选；反之会丢失弱目标。", "设置五组候选参数，综合数量稳定性、眼睛证据和耗时初筛，再结合对比图人工核验。", "当前单人大脸样例选中 1.20、6、50×50 的较严格组合，结果稳定且速度较快。"),
        ("GUI 与算法模块耦合风险", "如果把预处理和检测逻辑直接写在按钮回调中，代码难以测试和复用。", "GUI 只读取参数并调用 src 模块，命令行与 GUI 共用同一套算法和配置对象。", "核心算法可独立单元测试，GUI 导入和启动也通过验证。"),
    ]
    for i, (phenomenon, cause, solution, effect) in enumerate(problems, 1):
        add_heading(doc, f"5.{i} {phenomenon}", 2)
        add_body(doc, f"问题现象：{phenomenon}。", bold_prefix="问题现象：")
        add_body(doc, f"问题原因：{cause}", bold_prefix="问题原因：")
        add_body(doc, f"解决方案：{solution}", bold_prefix="解决方案：")
        add_body(doc, f"最终效果：{effect}", bold_prefix="最终效果：")

    add_heading(doc, "6 课程设计总结与心得", 1)
    add_heading(doc, "6.1 设计总结", 2)
    add_body(doc, "本课程设计完成了一个可运行的人脸图像预处理及检测系统。系统提供 PyQt6 图形界面和命令行两种入口，支持单图、批量、参数传入和参数搜索；预处理部分自主实现灰度化、Gamma、直方图均衡、高斯核、自定义卷积、高斯滤波和中值滤波，并封装 CLAHE 与尺寸归一化；检测部分加载 OpenCV Haar XML 完成人脸和人眼检测；输出部分保存检测图、对比图、CSV 与 JSON 指标。")
    add_body(doc, "系统的优势是模块边界清晰、中文路径安全、实验结果可复现，并且没有将无标注稳定性包装成准确率。不足之处也较明显：Haar 正面模型对大角度侧脸、遮挡和极小人脸能力有限；测试集只有三张控制变量演示图，没有人工标注框，因此不能严格计算 Precision、Recall 或 IoU；复杂真实场景中的鲁棒性仍需扩大样本后验证。后续可以增加侧脸级联、摄像头实时输入和基于深度学习的检测器，并制作人工标注测试集。")

    add_heading(doc, "6.2 心得体会", 2)
    add_body(doc, "这次课程设计让我对“预处理不是越多越好”有了更直观的认识。以前只知道直方图均衡和滤波的公式，真正把它们放到检测流程里后才发现，全局均衡可能把背景和噪声一起增强，高斯滤波虽然能抑制随机噪声，却也会损失边缘。因此，每种处理都应该针对具体问题，并保留对照分支验证效果。")
    add_body(doc, "在代码实现方面，我把灰度化、Gamma、均衡化和滤波拆成可单独测试的函数，再让命令行和 GUI 调用同一套模块。这个过程比写一个能跑的脚本更麻烦，但后期调试明显更轻松。参数搜索也让我体会到速度、漏检和误检之间没有绝对最优值，必须结合数据集和实际使用场景折中。后续如果继续完善，我希望加入人工标注和深度学习检测方法，用更客观的指标与 Haar 方法进行比较。")

    reference_heading = add_heading(doc, "参考文献", 1)
    reference_heading.paragraph_format.page_break_before = True
    references = [
        "[1] 冈萨雷斯, 伍兹. 数字图像处理（第四版）[M]. 北京: 电子工业出版社, 2020.",
        "[2] OpenCV. OpenCV Documentation[EB/OL]. https://docs.opencv.org/4.x/.",
        "[3] Viola P, Jones M. Rapid Object Detection using a Boosted Cascade of Simple Features[C]. CVPR, 2001.",
        "[4] Viola P, Jones M. Robust Real-Time Face Detection[J]. International Journal of Computer Vision, 2004, 57(2): 137-154.",
        "[5] Bradski G. The OpenCV Library[J]. Dr. Dobb's Journal of Software Tools, 2000.",
        "[6] Zuiderveld K. Contrast Limited Adaptive Histogram Equalization[A]. Graphics Gems IV[C]. Academic Press, 1994.",
        "[7] Riverbank Computing. PyQt6 Documentation[EB/OL]. https://www.riverbankcomputing.com/static/Docs/PyQt6/.",
    ]
    for ref in references:
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(ref)
        set_run_font(r, size=10.5)

    # 请求 Word 打开时更新目录和页码；随后还会通过 Word COM 主动更新。
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
